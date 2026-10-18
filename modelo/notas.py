"""
Módulo 3 · Calificaciones → rendimiento por área.

Reproduce el notebook (celdas 3.3 a 3.11 y la parte de notas del 9.1):
  3.3 solo secundaria (9.º–12.º) · 3.4 clasificación por reglas · 3.5 escala → 0–10
  3.6 promedio de semestres · 3.7 z contra la cohorte (materia×grado si n ≥ 5, si no el grado)
  3.8 agregación materia → área por año · 3.9 combinación temporal (γ = 0.8 desde el último grado)
"""
from __future__ import annotations

import csv
import io
import math
import re
from collections import defaultdict

import numpy as np

from .configuracion import Regla, clasificar
from .constantes import CODIGOS, GAMMA_ANIOS, GRADOS_SECUNDARIA, N_MIN_MATERIA
from .estadistica import corr_pareada, media_sd, phi

# Columnas del export del sistema académico (notebook 3.1)
COL = {"est": "Student ID (System)", "grado": "Grade Level 1", "anio": "Year", "materia": "Course",
       "nota": "Final grade AVG", "letra": "Final grade", "grado_actual": "Grade Level"}
OBLIGATORIAS = ["est", "grado", "anio", "materia", "nota"]


# ───────────────────────── lectura del reporte ─────────────────────────
def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None if isinstance(v, float) and math.isnan(v) else float(v)
    t = str(v).strip().replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def _filas_de_archivo(contenido: bytes, nombre: str):
    if nombre.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        cab = [str(c).strip() if c is not None else "" for c in next(it)]
        return [dict(zip(cab, fila)) for fila in it]
    texto = contenido.decode("utf-8-sig", errors="replace")
    muestra = texto[:4096]
    sep = ";" if muestra.count(";") > muestra.count(",") else ","
    return [{(k or "").strip(): v for k, v in r.items()} for r in csv.DictReader(io.StringIO(texto), delimiter=sep)]


def leer_reporte(contenido: bytes, nombre: str = "reporte.csv"):
    """Export del colegio (CSV o Excel) → lista de registros normalizados. Lanza ValueError si faltan columnas."""
    filas = _filas_de_archivo(contenido, nombre)
    if not filas:
        raise ValueError("El archivo está vacío.")
    faltan = [COL[c] for c in OBLIGATORIAS if COL[c] not in filas[0]]
    if faltan:
        raise ValueError("Al reporte le faltan columnas: " + ", ".join(faltan))
    out = []
    for r in filas:
        sid = str(r.get(COL["est"]) or "").strip()
        if sid.endswith(".0"):
            sid = sid[:-2]
        if not sid:
            continue
        g = _num(r.get(COL["grado"]))
        out.append({"id_colegio": sid, "grado_actual": _num(r.get(COL["grado_actual"])),
                    "grado": int(g) if g is not None and float(g).is_integer() else None,
                    "anio": str(r.get(COL["anio"]) or "").strip(), "materia": str(r.get(COL["materia"]) or ""),
                    "nota": _num(r.get(COL["nota"])), "letra": str(r.get(COL["letra"]) or "").strip()})
    return out


# ───────────────────────── 3.5 escala ─────────────────────────
def a_0_10(v, letra):
    """→ (nota en 0–10 o None, escala detectada). Un 0 con letra aprobatoria es 'sin calificación'."""
    if v is None:
        return None, "vacía"
    letra = (letra or "").strip().upper()
    if v == 0 and letra[:1] in ("A", "B", "C") and letra:
        return None, "sin calificación (0 con letra aprobatoria)"
    if v > 10:
        return v / 10, "0-100"
    if v > 4:
        return v, "0-10"
    return v * 2.5, "0-4"


def periodo_y_semestre(anio: str):
    m = re.search(r"(\d{4}\s*-\s*\d{4})", anio or "")
    s = re.search(r"Sem\s*(\d)", anio or "")
    return (m.group(1) if m else None), (s.group(1) if s else None)


# ───────────────────────── 3.3–3.6 preparación por estudiante ─────────────────────────
def preparar(registros: list[dict], reglas: list[Regla]):
    """Registros de UN estudiante → dict con filas clasificadas (9.º–12.º), promedio por materia y grado actual."""
    filas, grupos, grado_actual = [], defaultdict(list), None
    for r in registros:
        if r.get("grado_actual") is not None:
            grado_actual = max(grado_actual or 0, int(r["grado_actual"]))
        g = r.get("grado")
        if g not in GRADOS_SECUNDARIA:
            continue
        regla, estado = clasificar(r["materia"], reglas)
        nota10, _ = a_0_10(r.get("nota"), r.get("letra"))
        periodo, sem = periodo_y_semestre(r.get("anio", ""))
        filas.append({"grado": g, "periodo": periodo, "sem": sem, "materia": r["materia"], "nota10": nota10,
                      "estado": estado, "regla": regla})
        if estado == "incluida" and nota10 is not None:
            grupos[(g, r["materia"])].append(nota10)
    reglas_de = {(f["grado"], f["materia"]): f["regla"] for f in filas}
    por_materia = {k: {"nota10": sum(v) / len(v), "regla": reglas_de[k]} for k, v in grupos.items()}
    return {"filas": filas, "por_materia": por_materia, "grado_actual": grado_actual}


# ───────────────────────── 3.7 normas ─────────────────────────
def ajustar_normas_materias(por_materia_cohorte: list[dict]) -> dict:
    """Media, sd y n por materia×grado y por grado (sobre los registros estudiante×grado×materia de la cohorte)."""
    mat, gra = defaultdict(list), defaultdict(list)
    for pm in por_materia_cohorte:
        for (g, materia), v in pm.items():
            mat[f"{g}|{materia}"].append(v["nota10"]); gra[str(g)].append(v["nota10"])
    return {"materia": {k: list(media_sd(v)) for k, v in mat.items()},
            "grado": {k: list(media_sd(v)) for k, v in gra.items()}}


def z_materia(nota10: float, grado: int, materia: str, normas: dict):
    m, sd, n = normas["materia"].get(f"{grado}|{materia}", (None, None, 0))
    if not n or n < N_MIN_MATERIA:
        m, sd, n = normas["grado"].get(str(grado), (None, None, 0))
    if m is None or not sd or sd <= 0:
        return None
    return (nota10 - m) / sd


# ───────────────────────── 3.8–3.9 rendimiento ─────────────────────────
def rendimiento(por_materia: dict, normas: dict):
    """→ (aporte, por_anio, rend). Áreas sin evidencia quedan sin entrada (NaN en el notebook, nunca cero)."""
    aporte = []
    for (g, materia) in sorted(por_materia, key=lambda k: (k[0], k[1])):
        v = por_materia[(g, materia)]
        z = z_materia(v["nota10"], g, materia, normas)
        if z is None:
            continue
        for a in CODIGOS:
            w = v["regla"].pesos.get(a, 0) if v["regla"] else 0
            if w > 0:
                aporte.append({"grado": g, "area": a, "materia": materia, "w": w, "z": z, "nota10": v["nota10"]})
    acc = defaultdict(lambda: {"wz": 0.0, "W": 0.0, "materias": set()})
    for x in aporte:
        d = acc[(x["area"], x["grado"])]
        d["wz"] += x["w"] * x["z"]; d["W"] += x["w"]; d["materias"].add(x["materia"])
    por_anio = {k: {"R": d["wz"] / d["W"], "W": d["W"], "n_materias": len(d["materias"])} for k, d in acc.items()}
    rend = {}
    if por_anio:
        T = max(g for (_, g) in por_anio)
        for a in CODIGOS:
            anios = [(g, d) for (aa, g), d in por_anio.items() if aa == a]
            if not anios:
                continue
            pesos = [(GAMMA_ANIOS ** (T - g)) * d["W"] for g, d in anios]
            sp = sum(pesos)
            rend[a] = {"R": sum(p * d["R"] for p, (_, d) in zip(pesos, anios)) / sp,
                       "W_total": sum(d["W"] for _, d in anios), "W_efectiva": sp,
                       "anios": len(anios), "n_materias": sum(d["n_materias"] for _, d in anios), "T": T}
    return aporte, por_anio, rend


def ajustar_normas_area(rends: list[dict]) -> dict:
    """z_notas del Módulo 5.1: R estandarizado por área en la cohorte."""
    return {a: list(media_sd([r[a]["R"] for r in rends if a in r]))[:2] for a in CODIGOS}


def ajustar_normas_trayectoria(por_anios: list[dict]) -> dict:
    acc = defaultdict(list)
    for pa in por_anios:
        for (a, g), d in pa.items():
            acc[f"{a}|{g}"].append(d["R"])
    return {k: list(media_sd(v))[:2] for k, v in acc.items()}


def z_notas(rend: dict, normas_area: dict) -> dict:
    out = {}
    for a in CODIGOS:
        if a in rend and normas_area.get(a) and normas_area[a][1]:
            m, sd = normas_area[a]
            out[a] = (rend[a]["R"] - m) / sd
    return out


def trayectoria(por_anio: dict, normas_tray: dict) -> list[dict]:
    out = []
    for (a, g) in sorted(por_anio, key=lambda k: (k[1], k[0])):
        n = normas_tray.get(f"{a}|{g}")
        if not n or not n[1]:
            continue
        z = (por_anio[(a, g)]["R"] - n[0]) / n[1]
        out.append({"grado": g, "area": a, "rendimiento_percentil": float(np.round(phi(z) * 100))})
    return out


def confiabilidad_notas(por_anios: dict, ids=None) -> dict:
    """α de las notas por área: consistencia entre años con Spearman-Brown, recortada a [0.30, 0.95]."""
    ids = list(por_anios) if ids is None else [i for i in por_anios if i in ids]
    out = {}
    for a in CODIGOS:
        grados = sorted({g for i in ids for (aa, g) in por_anios[i] if aa == a})
        if len(grados) < 2:
            out[a] = 0.30
            continue
        M = np.array([[por_anios[i][(a, g)]["R"] if (a, g) in por_anios[i] else np.nan for g in grados] for i in ids
                      if any((a, g) in por_anios[i] for g in grados)])
        cs = [corr_pareada(M[:, x], M[:, y]) for x in range(len(grados)) for y in range(x + 1, len(grados))]
        r = float(np.nanmean(cs)) if not all(math.isnan(c) for c in cs) else float("nan")
        k = len(grados)
        sb = k * r / (1 + (k - 1) * r)
        out[a] = float(np.clip(sb, 0.30, 0.95)) if not math.isnan(sb) else 0.30
    return out


# ───────────────────────── contrato web ─────────────────────────
def materias_clave(aporte: list[dict], area: str, T: int, top: int = 3):
    """Materias del último grado que más pesan en el área (orden: peso, nota; empate → nombre)."""
    ult = [x for x in aporte if x["area"] == area and x["grado"] == T]
    ult.sort(key=lambda x: x["materia"])
    ult.sort(key=lambda x: (-x["w"], -(x["nota10"] if x["nota10"] is not None else -1e9)))
    return [{"materia": x["materia"], "nota": x["nota10"], "peso": x["w"]} for x in ult[:top]]


def origen_notas(filas: list[dict]) -> list[dict]:
    """Detalle de notas para la sección «Fuentes de datos» de la interfaz (notebook 9.1 notas_origen)."""
    grupos = defaultdict(list)
    for f in filas:
        grupos[(f["grado"], f["periodo"])].append(f)
    anios = []
    for (g, per) in sorted(grupos, key=lambda k: (k[0], k[1] is None, k[1] or "")):
        por_mat = defaultdict(list)
        for f in grupos[(g, per)]:
            por_mat[f["materia"]].append(f)
        mats = []
        for m in sorted(por_mat):
            k = por_mat[m]
            s1 = [x["nota10"] for x in k if x["sem"] == "1" and x["nota10"] is not None]
            s2 = [x["nota10"] for x in k if x["sem"] == "2" and x["nota10"] is not None]
            todas = [x["nota10"] for x in k if x["nota10"] is not None]
            regla = k[0]["regla"]
            mats.append({"materia": m, "sem1": sum(s1) / len(s1) if s1 else None, "sem2": sum(s2) / len(s2) if s2 else None,
                         "final": sum(todas) / len(todas) if todas else None, "estado": k[0]["estado"],
                         "areas": regla.areas if regla else ""})
        orden = {"incluida": 0, "sin_regla": 1, "excluida": 2}
        mats.sort(key=lambda x: (orden.get(x["estado"], 3), x["materia"]))
        anios.append({"grado": int(g), "periodo": per, "materias": mats})
    return anios
